"""Экспорт данных в Excel (файл собирается в памяти и отправляется документом).

Два отчёта:
- `metric_rows_to_xlsx` — выгрузка одного показателя (период/значение/источник);
- `expenses_report_xlsx` — отчёт по личным расходам из бюджета компании: строки —
  месяцы, колонка «Личные расходы», разбивка по категориям отдельным листом.

Отчёт по расходам собирается ИЗ БАЗЫ каждый раз, а не дописывается в
загруженный файл компании: оригиналы документов остаются нетронутыми, а цифры
в отчёте всегда соответствуют тому, что записано в журнале.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .money import to_decimal

MONTHS_NOM = (
    "", "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
)

_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_MONEY_FORMAT = "#,##0.00"


def metric_rows_to_xlsx(title: str, rows: list[dict[str, Any]]) -> bytes:
    """rows — результат storage.series_for_metric: period_label/value/unit/
    currency/document_name/sheet/cell_ref."""
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Показатель")[:31]
    ws.append(["Период", "Значение", "Единица", "Валюта", "Документ", "Лист", "Ячейка"])
    for r in rows:
        ws.append([
            r.get("period_label"),
            r.get("value"),
            r.get("unit"),
            r.get("currency"),
            r.get("document_name"),
            r.get("sheet"),
            r.get("cell_ref"),
        ])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["E"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _as_decimal(value: Any) -> Decimal:
    if isinstance(value, Decimal):
        return value
    return to_decimal(value) or Decimal(0)


def _month_key(when: date) -> tuple[int, int, str]:
    return when.year, when.month, f"{when.month:02d}.{when.year}"


def expenses_report_xlsx(
    rows: list[dict[str, Any]],
    *,
    org_name: str = "",
    today: date | None = None,
) -> bytes:
    """Отчёт по личным расходам: месяц × категория, с колонкой «Личные расходы».

    rows — результат `app.expenses.ledger_rows` (when/kind/amount/category/description).
    Лист «Отчёт»: по строке на месяц — расходы, доходы, сальдо, число записей.
    Лист «По категориям»: категории расходов по строкам, месяцы по колонкам.
    """
    today = today or date.today()
    # месяц -> категория -> сумма (только расходы); месяц -> доходы
    by_month: dict[tuple[int, int], dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
    incomes: dict[tuple[int, int], Decimal] = defaultdict(Decimal)
    counts: dict[tuple[int, int], int] = defaultdict(int)

    for r in rows:
        when = r.get("when")
        if not isinstance(when, date):
            continue
        year, month, _label = _month_key(when)
        amount = _as_decimal(r.get("amount"))
        counts[(year, month)] += 1
        if r.get("kind") == "income":
            incomes[(year, month)] += amount
        else:
            category = str(r.get("category") or "Прочее")
            by_month[(year, month)][category] += amount

    months = sorted(set(by_month) | set(incomes))
    wb = Workbook()

    # --- лист 1: месяцы, колонка «Личные расходы» + топ категорий ---
    ws = wb.active
    ws.title = "Отчёт"
    title = "Личные расходы из бюджета компании"
    if org_name:
        title += f" — {org_name}"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Сформировано {today:%d.%m.%Y} · всего записей: {len(rows)}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(["Месяц", "Личные расходы", "Личные доходы", "Сальдо", "Записей"])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    total_expense = total_income = Decimal(0)
    for year, month in months:
        cats = by_month.get((year, month), {})
        expense = sum(cats.values(), Decimal(0))
        income = incomes.get((year, month), Decimal(0))
        total_expense += expense
        total_income += income
        ws.append([
            f"{MONTHS_NOM[month]} {year}",
            float(expense),
            float(income),
            float(expense - income),
            counts.get((year, month), 0),
        ])

    ws.append([
        "ИТОГО",
        float(total_expense),
        float(total_income),
        float(total_expense - total_income),
        sum(counts.values()),
    ])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    # крупнейшие категории расходов — сразу под таблицей, чтобы было видно
    # структуру, не переходя на второй лист
    all_cats: dict[str, Decimal] = defaultdict(Decimal)
    for cats in by_month.values():
        for name, value in cats.items():
            all_cats[name] += value
    if all_cats:
        ws.append([])
        ws.append(["Структура расходов по категориям"])
        ws[ws.max_row][0].font = Font(bold=True)
        for name, value in sorted(all_cats.items(), key=lambda kv: -kv[1]):
            share = f"{value / total_expense * 100:.1f}%" if total_expense else "—"
            ws.append([name, float(value), share])

    ws.column_dimensions["A"].width = 34
    for col in "BC":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 10
    for row in ws.iter_rows(min_row=header_row + 1, min_col=2, max_col=4):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = _MONEY_FORMAT

    # --- лист 2: категории по месяцам ---
    ws2 = wb.create_sheet("По категориям")
    if months and all_cats:
        month_labels = [f"{MONTHS_NOM[m]} {y}" for y, m in months]
        ws2.append(["Категория", *month_labels, "Итого"])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
        for name, _total in sorted(all_cats.items(), key=lambda kv: -kv[1]):
            values = [float(by_month.get((y, m), {}).get(name, Decimal(0))) for y, m in months]
            ws2.append([name, *values, sum(values)])
        ws2.append(["ИТОГО", *[float(sum(by_month.get(k, {}).values(), Decimal(0))) for k in months],
                    float(total_expense)])
        for cell in ws2[ws2.max_row]:
            cell.font = Font(bold=True)
        for row in ws2.iter_rows(min_row=2, min_col=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = _MONEY_FORMAT
        ws2.column_dimensions["A"].width = 26
        for idx in range(2, len(months) + 3):
            ws2.column_dimensions[get_column_letter(idx)].width = 15
    else:
        ws2.append(["Нет записей журнала — расходы ещё не вносились."])

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def expenses_report_filename(today: date | None = None) -> str:
    return f"Личные_расходы_{(today or date.today()):%Y-%m}.xlsx"
