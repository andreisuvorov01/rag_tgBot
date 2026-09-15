"""Тесты усиления векторной базы (offline-«живость»): расширение запросов,
мультивариантный гибридный поиск, эвристический реранкер, карточки показателей,
резолв по аббревиатурам."""
import openpyxl

from app.config import settings
from app.embeddings import EmbeddingService
from app.ingest.pipeline import process_document
from app.llm import make_llm
from app.qa import AnswerPipeline
from app.query_expansion import expand_query
from app.rag import hybrid_search
from app.rerank import HeuristicReranker, make_reranker
from app.storage import make_engine, make_sessionmaker

# ------------------------------------------------------------ расширение запросов

def test_expand_query_abbreviations():
    variants = expand_query("сколько ндс в выписке")
    assert len(variants) == 2
    assert "налог на добавленную стоимость" in variants[1]
    assert variants[0] == "сколько ндс в выписке"


def test_expand_query_synonyms():
    variants = expand_query("как изменилась дебиторка за год")
    assert any("дебиторская задолженность" in v for v in variants)


def test_expand_query_no_change():
    # «риски» нет ни в группах синонимов, ни в сокращениях — вариант один
    assert expand_query("какие риски упоминаются?") == ["какие риски упоминаются?"]
    assert expand_query("") == [""]


def test_provider_field_with_model_name_self_corrects():
    """Частая ошибка: имя модели в EMBEDDINGS_PROVIDER — трактуется как
    sentence_transformers + модель. На машине без ST — откат на hash."""
    s = settings.model_copy(update={"embeddings_provider": "BAAI/bge-m3"})
    if settings._st_available():
        assert s.effective_embeddings_provider == "sentence_transformers"
        assert s.effective_embeddings_model == "BAAI/bge-m3"
    else:
        assert s.effective_embeddings_provider == "hash"


# ------------------------------------------------------------ эвристический реранкер

def test_heuristic_reranker_prefers_relevant():
    r = HeuristicReranker(settings)
    chunks = [
        {"id": 1, "body": "Общие слова о деятельности компании.", "semantic_score": 0.1},
        {"id": 2, "body": "Дебиторская задолженность выросла на 15% за год.", "semantic_score": 0.6},
    ]
    import asyncio

    result = asyncio.run(r.rerank("дебиторская задолженность", chunks, k=2))
    assert result[0]["id"] == 2


def test_make_reranker_auto_mock_uses_heuristic():
    s = settings.model_copy(update={"reranker": "auto", "llm_provider": "mock"})
    assert isinstance(make_reranker(s, make_llm(s)), HeuristicReranker)


# ------------------------------------------------------------ поиск по синонимам

async def test_hybrid_search_finds_by_synonym(tmp_path):
    """«дебиторка» находит фрагмент про «дебиторскую задолженность» без API."""
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    settings.reranker = "heuristic"
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)

    import io

    import docx as docx_lib

    d = docx_lib.Document()
    d.add_paragraph("Дебиторская задолженность покупателей увеличилась на 15% к концу года.")
    d.add_paragraph("Расходы на аренду офиса остались на уровне прошлого года.")
    buf = io.BytesIO()
    d.save(buf)
    async with sessions() as s:
        await process_document(s, emb, settings, org_id=1, user_id=1,
                               original_name="notes.docx", content=buf.getvalue())
        await s.commit()
        found = await hybrid_search(s, emb, 1, "что с дебиторкой?", k=3, user_id=1)
    assert found
    assert "дебиторская задолженность" in found[0]["body"].casefold()
    assert all("rrf_score" in f for f in found)
    await llm.close()
    await engine.dispose()


# ------------------------------------------------------------ карточки показателей

async def test_metric_cards_created_and_unique(tmp_path):
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Показатель", "2023", "2024"])
    ws.append(["Дебиторская задолженность", "1000", "1200"])
    src = tmp_path / "c.xlsx"
    wb.save(src)
    async with sessions() as s:
        await process_document(s, emb, settings, org_id=1, user_id=1,
                               original_name="c.xlsx", content=src.read_bytes())
        await s.commit()

        from sqlalchemy import select

        from app.storage import Chunk

        cards = (await s.scalars(
            select(Chunk).where(Chunk.section.like("карточка:%"))
        )).all()
        assert len(cards) == 1
        assert "2023" in cards[0].body and "2024" in cards[0].body
        assert "дебиторская задолженность" in cards[0].body

        # перечитывание не создаёт вторую карточку
        from app.ingest.pipeline import reparse_document
        from app.storage import Document

        doc_row = (await s.scalars(select(Document))).first()
        await reparse_document(s, emb, settings, document_id=doc_row.id, user_id=1)
        await s.commit()
        cards_after = (await s.scalars(
            select(Chunk).where(Chunk.section.like("карточка:%"))
        )).all()
        assert len(cards_after) == 1

    # поиск по карточке через синоним
    async with sessions() as s:
        found = await hybrid_search(s, emb, 1, "дебиторка", k=3, user_id=1)
    assert any("Карточка показателя" in f["body"] for f in found)

    await llm.close()
    await engine.dispose()


async def test_resolve_metric_by_abbreviation(tmp_path):
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)
    pipeline = AnswerPipeline(sessions, emb, llm, settings, org_names=[])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Показатель", "2024"])
    ws.append(["Налог на добавленную стоимость", "250"])
    src = tmp_path / "nds.xlsx"
    wb.save(src)
    async with sessions() as s:
        await process_document(s, emb, settings, org_id=1, user_id=1,
                               original_name="nds.xlsx", content=src.read_bytes())
        await s.commit()

    outcome = await pipeline.answer(1, 1, "Сколько составляет НДС за 2024?")
    assert "250" in outcome.text
    await llm.close()
    await engine.dispose()
