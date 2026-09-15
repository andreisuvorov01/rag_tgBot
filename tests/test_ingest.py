"""Тесты извлечения таблиц из Excel/CSV и вертикального макета."""
import csv

import openpyxl

from app.ingest.load_excel import load_csv, load_excel


def _make_wide(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Опер. расходы"
    ws["A1"] = "Отчет за 2025 год"
    ws["A2"] = "тыс. руб."
    ws.append(["Показатель", "2024", "2025"])
    ws.append(["Выручка", "41 200", "45 800"])
    ws.append(["Аренда спецтехники", "9 800", "11 400"])
    ws.append(["Прочие расходы", "(500)", "—"])
    wb.save(path)


def test_load_excel_wide(tmp_path):
    p = tmp_path / "r.xlsx"
    _make_wide(p)
    doc = load_excel(str(p))
    assert doc.warnings == [] or all("скан" not in w for w in doc.warnings)
    facts = {(f.metric_name, f.period.label): f for f in doc.facts}
    v = facts[("выручка", "2025")]
    assert v.value == 45_800_000.0  # множитель "тыс. руб." применён
    assert v.currency == "RUB"
    assert facts[("аренда спецтехники", "2024")].value == 9_800_000.0
    assert facts[("прочие расходы", "2024")].value == -500_000.0
    assert ("прочие расходы", "2025") not in facts  # "—" -> нет значения
    assert doc.chunks  # лист сохранён как текстовый чанк для RAG


def test_load_excel_vertical(tmp_path):
    p = tmp_path / "v.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "По регионам"
    ws.append(["Период", "Москва", "Регионы"])
    ws.append(["2023", "5 600", "4 200"])
    ws.append(["2024", "6 900", "4 500"])
    ws.append(["2025", "7 800", "4 900"])
    wb.save(p)
    doc = load_excel(str(p))
    facts = {(f.metric_name, f.period.label): f for f in doc.facts}
    assert facts[("москва", "2024")].value == 6900.0
    assert facts[("регионы", "2025")].value == 4900.0


def test_load_csv(tmp_path):
    p = tmp_path / "d.csv"
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Показатель", "2023", "2024"])
        w.writerow(["Выручка", "10 000", "12 000"])
        w.writerow(["Расходы", "8 000", "9 500"])
    doc = load_csv(str(p))
    facts = {(f.metric_name, f.period.label): f for f in doc.facts}
    assert facts[("выручка", "2024")].value == 12000.0
    assert facts[("расходы", "2023")].value == 8000.0


async def test_duplicate_report_carries_original_stats(tmp_path):
    """Повторная загрузка того же файла: сводка несёт реальные данные
    первой версии, а не нули нового вызова (прод-кейс 13.09.2026)."""
    import openpyxl

    from app.config import settings
    from app.embeddings import EmbeddingService
    from app.ingest.pipeline import process_document
    from app.storage import make_engine, make_sessionmaker

    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Показатель", "2024"])
    ws.append(["Выручка", "100"])
    ws.append(["Аренда", "50"])
    src = tmp_path / "doc.xlsx"
    wb.save(src)

    async with sessions() as s:
        r1 = await process_document(s, emb, settings, org_id=1, user_id=1,
                                    original_name="doc.xlsx", content=src.read_bytes())
        await s.commit()
    assert r1.status == "processed" and r1.facts == 2

    async with sessions() as s:
        r2 = await process_document(s, emb, settings, org_id=1, user_id=1,
                                    original_name="doc.xlsx", content=src.read_bytes())
        await s.commit()

    assert r2.status == "duplicate"
    assert r2.facts == 2 and r2.chunks == r1.chunks
    assert r2.periods == r1.periods
    summary = r2.summary()
    assert "уже загружен и обработан" in summary
    assert "Показателей: 2" in summary
    assert "2024" in summary
    await engine.dispose()


async def test_zero_facts_hint_for_text_only_sheet(tmp_path):
    """Таблиц с периодами нет, но текст есть: фактов 0, текст проиндексирован,
    предупреждение объясняет, что произошло."""
    import openpyxl

    from app.config import settings
    from app.embeddings import EmbeddingService
    from app.ingest.pipeline import process_document
    from app.storage import make_engine, make_sessionmaker

    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Комментарий к отчёту без таблиц"])
    ws.append(["Выручка выросла за счёт новых направлений."])
    src = tmp_path / "text.xlsx"
    wb.save(src)

    async with sessions() as s:
        report = await process_document(s, emb, settings, org_id=1, user_id=1,
                                        original_name="text.xlsx", content=src.read_bytes())
        await s.commit()
    assert report.status == "processed"
    assert report.facts == 0 and report.chunks > 0
    assert any("колонками периодов" in w for w in report.warnings)
    await engine.dispose()
