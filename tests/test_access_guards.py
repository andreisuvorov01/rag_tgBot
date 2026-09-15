"""Защита границ доверия: Text-to-SQL guard, изоляция организаций в callback-
действиях, авторизация инлайн-кнопок."""
import openpyxl

from app.bot import BotApp, CallbackAuth
from app.config import settings
from app.embeddings import EmbeddingService
from app.ingest.pipeline import process_document, reparse_document
from app.llm import make_llm
from app.qa import AnswerPipeline, sql_is_safe
from app.storage import confirm_facts, make_engine, make_sessionmaker, register_user, supersedes
from app.summarize import document_summary
from tests.test_bot_flows import FakeCallback, FakeMessage


def test_sql_guard_accepts_legitimate_query():
    assert sql_is_safe(
        "SELECT m.name, SUM(f.value) FROM facts_visible f JOIN metrics m ON m.id = f.metric_id "
        "WHERE EXTRACT(YEAR FROM f.period_end) = 2025 GROUP BY m.name"
    )


def test_sql_guard_rejects_foreign_tables_and_functions():
    bad = [
        "SELECT telegram_id FROM users WHERE 'facts_visible' = 'facts_visible'",
        "SELECT * FROM facts_visible f JOIN documents d ON d.id = f.document_id",
        "SELECT * FROM facts_visible WHERE metric_id IN (SELECT user_id FROM audit_log)",
        "SELECT pg_read_file('/etc/passwd') FROM facts_visible",
        "SELECT pg_sleep(10) FROM facts_visible",
        "SELECT * FROM facts_visible; DROP TABLE facts",
        "SELECT * FROM facts",
        "UPDATE facts SET value = 0",
    ]
    for sql in bad:
        assert not sql_is_safe(sql), sql


async def _two_orgs(tmp_path):
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)
    pipeline = AnswerPipeline(sessions, emb, llm, settings, org_names=[])
    async with sessions() as s:
        await register_user(s, 1, "A", "ООО Альфа")   # org 1
        await register_user(s, 2, "B", "ООО Бета")    # org 2
        await s.commit()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Показатель", "2023", "2024"])
    ws.append(["Выручка", "100", "120"])
    src = tmp_path / "r.xlsx"
    wb.save(src)
    async with sessions() as s:
        report = await process_document(s, emb, settings, org_id=1, user_id=1,
                                        original_name="r.xlsx", content=src.read_bytes())
        await s.commit()
    return engine, sessions, emb, pipeline, report.document_id


async def test_callback_actions_isolated_by_org(tmp_path):
    engine, sessions, emb, pipeline, doc_id = await _two_orgs(tmp_path)
    async with sessions() as s:
        # пользователь org 2 подделал callback с чужим document_id
        assert await confirm_facts(s, doc_id, org_id=2) == 0
        assert await supersedes(s, doc_id, doc_id + 1, org_id=2) is None
        assert (await reparse_document(s, emb, settings, document_id=doc_id, user_id=2, org_id=2)).status == "failed"
        assert await document_summary(pipeline, s, 2, doc_id, 2) == "Документ не найден"
        # владелец — всё доступно
        assert await confirm_facts(s, doc_id, org_id=1) >= 1
        assert (await reparse_document(s, emb, settings, document_id=doc_id, user_id=1, org_id=1)).status == "processed"
    await engine.dispose()


async def test_callback_auth_middleware_blocks_unregistered(tmp_path):
    engine, sessions, emb, pipeline, _ = await _two_orgs(tmp_path)
    app = BotApp(settings, pipeline, sessions)
    mw = CallbackAuth(app)
    called = []

    async def handler(event, data):
        called.append(event.from_user.id)

    await mw(handler, FakeCallback(FakeMessage(), "menu:metrics", user_id=999), {})
    await mw(handler, FakeCallback(FakeMessage(), "menu:metrics", user_id=1), {})
    assert called == [1]
    # org_id берётся из call.from_user, а не из call.message (там автор — бот)
    assert await app._org_id(FakeCallback(FakeMessage(user_id=42), "x", user_id=2)) == 2
    await engine.dispose()


async def test_clear_chat_history_batches_and_stops_at_48h_boundary():
    """Пачки по 100 вниз; отклонённая пачка добирается поштучно; пачка без
    удалений (граница 48 ч) завершает обход."""
    from aiogram.exceptions import TelegramBadRequest

    from app.bot import clear_chat_history

    class FakeBot:
        def __init__(self, deletable: set[int]):
            self.deletable = deletable
            self.bulk: list[list[int]] = []
            self.single: list[int] = []

        async def delete_messages(self, chat_id, ids):
            self.bulk.append(list(ids))
            if not set(ids) <= self.deletable:
                raise TelegramBadRequest(method=None, message="message can't be deleted")
            return True

        async def delete_message(self, chat_id, mid):
            self.single.append(mid)
            if mid in self.deletable:
                return True
            raise TelegramBadRequest(method=None, message="message can't be deleted")

    bot = FakeBot(deletable=set(range(151, 251)))  # 250 = /clear, старше 150 — >48 ч
    await clear_chat_history(bot, chat_id=1, last_message_id=250)
    assert bot.bulk[0] == list(range(250, 150, -1))          # первая пачка целиком удалена
    assert bot.bulk[1] == list(range(150, 50, -1))           # вторая отклонена -> поштучно
    assert bot.single == list(range(150, 50, -1)) and len(bot.bulk) == 2  # ничего не удалилось -> стоп


async def test_delete_user_documents_wipes_data_and_orphan_metrics(tmp_path):
    from pathlib import Path

    from sqlalchemy import func, select

    from app.storage import Chunk, Document, Fact, Metric, delete_user_documents

    engine, sessions, emb, pipeline, doc_id = await _two_orgs(tmp_path)
    async with sessions() as s:
        stored = Path((await s.get(Document, doc_id)).stored_path)
        assert stored.exists()
        # чужой пользователь той же схемы ничего не удаляет
        assert await delete_user_documents(s, org_id=2, user_id=2) == []
        paths = await delete_user_documents(s, org_id=1, user_id=1)
        await s.commit()
    assert paths == [str(stored)]
    async with sessions() as s:
        for model in (Document, Fact, Chunk, Metric):
            assert (await s.execute(select(func.count()).select_from(model))).scalar_one() == 0
    await engine.dispose()


async def test_upload_reports_progress_stages(tmp_path):
    """process_document зовёт progress по этапам, доля растёт монотонно до ~1."""
    from app.bot import ProgressBar
    from tests.test_bot_flows import FakeMessage

    engine, sessions, emb, pipeline, _ = await _two_orgs(tmp_path)
    src = tmp_path / "r.xlsx"
    status = FakeMessage(user_id=1)
    bar = ProgressBar(status, "📄 r.xlsx", min_interval=0.0)
    stages: list[tuple[str, float]] = []

    async def spy(stage, fraction):
        stages.append((stage, fraction))
        await bar(stage, fraction)

    async with sessions() as s:
        await process_document(s, emb, settings, org_id=2, user_id=2, original_name="r.xlsx",
                               content=src.read_bytes(), progress=spy)
    fractions = [f for _, f in stages]
    assert fractions == sorted(fractions) and fractions[-1] >= 0.97
    assert any("векторизация" in st for st, _ in stages)
    edited = [x["text"] for x in status.sent if x.get("edited")]
    assert edited and "▰▰▰▰▰▰▰▰▰▰ 97%" in edited[-1]
    await engine.dispose()


async def test_ack_swallows_expired_callback_query():
    """Долгий обработчик (перечитывание > 15 с) не должен падать на answerCallbackQuery."""
    from aiogram.exceptions import TelegramBadRequest

    from app.bot import ack

    class Expired:
        async def answer(self, text=None, show_alert=False):
            raise TelegramBadRequest(method=None, message="Bad Request: query is too old and response timeout expired or query ID is invalid")

    class Other:
        async def answer(self, text=None, show_alert=False):
            raise TelegramBadRequest(method=None, message="Bad Request: something else")

    await ack(Expired())  # не бросает
    import pytest
    with pytest.raises(TelegramBadRequest):
        await ack(Other())


def test_mock_classifier_routes_document_questions_to_text_search():
    from app.llm import _mock_classify

    for q in ("кто автор документа", "как называется организация", "какой инн у организации",
              "где находится компания", "что говорится об аудите"):
        assert _mock_classify(q)["intent"] == "explain", q
    assert _mock_classify("какие коммерческие расходы были в 2025")["metric_query"] == "коммерческие расходы"


def test_focus_snippet_picks_matching_lines():
    from app.rag import focus_snippet

    body = "Информация из ресурса БФО\nДата формирования | 13.09.2026\nИНН | 7721546864\n" \
           "Признак лица, подписавшего документ | представитель\nФИО руководителя | Лесных Юлия\nСтраница 1"
    out = focus_snippet(body, "кто руководитель", radius=0)
    assert out == "ФИО руководителя | Лесных Юлия"
    assert focus_snippet(body, "ничего общего", radius=0).startswith("Информация")


async def test_basic_rag_questions_offline(tmp_path):
    """Без внешней LLM: показатель — как подстрока вопроса, реквизиты — из текста,
    короткий новый вопрос не подхватывает показатель из памяти диалога."""
    engine, sessions, emb, pipeline, _ = await _two_orgs(tmp_path)
    from docx import Document as Docx

    d = Docx()
    d.add_paragraph("Полное наименование юридического лица | ООО Альфа")
    d.add_paragraph("Местонахождение (адрес) | г. Москва, ул. Ленина, 1")
    d.add_paragraph("ФИО руководителя | Иванов Иван Иванович")
    d.add_paragraph("Среднесписочная численность сотрудников | 120 человек")
    src = tmp_path / "info.docx"
    d.save(src)
    async with sessions() as s:
        await process_document(s, emb, settings, org_id=1, user_id=1, original_name="info.docx", content=src.read_bytes())
        await s.commit()
    out = (await pipeline.answer(1, 1, "какая выручка была в 2024 году")).text
    assert "120" in out and "выручка" in out
    assert "Иванов" in (await pipeline.answer(1, 1, "кто автор документа")).text
    assert "Ленина" in (await pipeline.answer(1, 1, "где находится компания")).text
    out = (await pipeline.answer(1, 1, "сколько сотрудников")).text
    assert "120 человек" in out and "Данные по показателю" not in out  # не факт о «выручке» из памяти
    await engine.dispose()
