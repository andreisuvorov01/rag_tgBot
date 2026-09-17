"""Тесты детальной работы с таблицами: многоярусные шапки с merged-ячейками,
план/факт, разделы и иерархия, «в т.ч.», единицы в ячейках, сверка итогов."""
from decimal import Decimal

import openpyxl

from app.config import settings
from app.embeddings import EmbeddingService
from app.ingest.load_excel import load_excel
from app.ingest.normalize import detect_variant, parse_number, parse_number_with_unit
from app.ingest.pipeline import process_document
from app.llm import make_llm
from app.storage import find_metric_by_name, make_engine, make_sessionmaker, series_for_metric

# ------------------------------------------------------------ нормализация

def test_number_formats():
    assert parse_number("1.234,56") == Decimal("1234.56")
    assert parse_number("2.500") == Decimal("2500")
    assert parse_number("12,5") == Decimal("12.5")
    assert parse_number("10.00") == Decimal("10.00")  # обычная точка-десятичная


def test_number_with_unit_in_cell():
    v, info = parse_number_with_unit("12,7 млн")
    assert v == Decimal("12700000") and info is not None
    v, info = parse_number_with_unit("9 800 тыс. руб.")
    assert v == Decimal("9800000") and info.currency == "RUB"
    v, info = parse_number_with_unit("(1 200) тыс.")
    assert v == Decimal("-1200000")
    assert parse_number_with_unit("Выручка")[0] is None


def test_detect_variant():
    assert detect_variant("2024 (план)") == "plan"
    assert detect_variant("план на 2025") == "plan"
    assert detect_variant("2024 факт") == "fact"
    assert detect_variant("бюджет") == "budget"
    assert detect_variant("прогноз") == "forecast"
    assert detect_variant("2024") is None


# ------------------------------------------------------------ парсер грида

def test_multilevel_header_plan_fact(tmp_path):
    """Шапка из двух ярусов с merged-ячейкой: 2024 (план | факт), 2025."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "План-факт"
    ws["A1"] = "тыс. руб."
    ws["B1"] = "2024"
    ws.merge_cells("B1:C1")
    ws["D1"] = "2025"
    ws.merge_cells("D1:D2")
    ws["A2"] = ""
    ws["B2"] = "план"
    ws["C2"] = "факт"
    ws["A3"] = "Выручка"
    ws["B3"] = "41 200"
    ws["C3"] = "41 950"
    ws["D3"] = "45 800"
    ws["A4"] = "Аренда"
    ws["B4"] = "9 800"
    ws["C4"] = "10 100"
    ws["D4"] = "11 400"
    p = tmp_path / "pf.xlsx"
    wb.save(p)

    doc = load_excel(str(p))
    got = {(f.metric_name, f.period.label, f.variant): f.value for f in doc.facts}
    assert got[("выручка", "2024", "plan")] == 41_200_000.0
    assert got[("выручка", "2024", "fact")] == 41_950_000.0
    assert got[("выручка", "2025", "fact")] == 45_800_000.0
    assert got[("аренда", "2024", "plan")] == 9_800_000.0
    # факт и план не слились в одну точку ряда
    assert ("выручка", "2024", "plan") != ("выручка", "2024", "fact")


def test_sections_totals_vth_units(tmp_path):
    """Раздел без чисел -> родитель; «Итого» с префиксом раздела;
    «в т.ч.» -> дочерний показатель; единица прямо в ячейке."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Расходы"
    ws["A1"] = "Показатель"
    ws["B1"] = "2025"
    ws["A2"] = "Расходы"  # строка-раздел: текст без чисел -> родительская группа
    ws["A3"] = "Зарплаты"
    ws["B3"] = "15 600"
    ws["A4"] = "Аренда"
    ws["B4"] = "3 300"
    ws["A5"] = "Итого"
    ws["B5"] = "18 900"
    ws["A6"] = "в т.ч. коммунальные"
    ws["B6"] = "2 950"
    ws["A7"] = "Прочее, млн руб."
    ws["B7"] = "1,2"
    p = tmp_path / "sec.xlsx"
    wb.save(p)

    doc = load_excel(str(p))
    got = {f.metric_name: f for f in doc.facts}
    assert "зарплаты" in got and "аренда" in got
    # итог получил префикс раздела и пометку
    total = got.get("расходы итого")
    assert total is not None and total.is_total and total.section == "расходы"
    # «в т.ч.» — дочерний к итогу
    child = got.get("расходы итого — коммунальные")
    assert child is not None and child.parent_hint == "расходы итого"
    # единица из названия строки уходит в unit/множитель, имя — чистое: 1,2 млн = 1 200 000
    other = got.get("прочее")
    assert other is not None and other.value == 1_200_000.0 and other.currency == "RUB"
    assert "расходы" in doc.sections


async def test_hierarchy_and_series_fact_only(tmp_path):
    """Интеграция: parent_id в словаре; ряд показывает только факт (не план)."""
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"] = "2024"
    ws.merge_cells("B1:C1")
    ws["D1"] = "2025"
    ws.merge_cells("D1:D2")
    ws["B2"] = "план"
    ws["C2"] = "факт"
    ws["A3"] = "Выручка"
    ws["B3"] = "40000"
    ws["C3"] = "41200"
    ws["D3"] = "45800"
    p = tmp_path / "h.xlsx"
    wb.save(p)

    async with sessions() as s:
        report = await process_document(s, emb, settings, org_id=1, user_id=1,
                                        original_name="h.xlsx", content=p.read_bytes())
        await s.commit()
    assert report.status == "processed"
    assert not any("не сходится" in w for w in report.warnings)

    async with sessions() as s:
        m = await find_metric_by_name(s, 1, "выручка")
        rows = await series_for_metric(s, 1, m.id)
    labels = {r["period_label"]: r["value"] for r in rows}
    assert labels == {"2024": 41_200.0, "2025": 45_800.0}  # план 40 000 не попал

    await llm.close()
    await engine.dispose()


async def test_totals_mismatch_warning(tmp_path):
    """Сумма строк не сходится с «Итого» -> предупреждение в отчёте загрузки."""
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Показатель", "2025"])
    ws.append(["Зарплаты", "500"])
    ws.append(["Аренда", "300"])
    ws.append(["Итого расходов", "999"])  # 500+300 != 999
    p = tmp_path / "m.xlsx"
    wb.save(p)

    async with sessions() as s:
        report = await process_document(s, emb, settings, org_id=1, user_id=1,
                                        original_name="m.xlsx", content=p.read_bytes())
        await s.commit()
    assert any("не сходится" in w for w in report.warnings)

    await llm.close()
    await engine.dispose()


async def test_totals_match_no_warning(tmp_path):
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Показатель", "2025"])
    ws.append(["Зарплаты", "500"])
    ws.append(["Аренда", "300"])
    ws.append(["Итого расходов", "800"])
    p = tmp_path / "m2.xlsx"
    wb.save(p)

    async with sessions() as s:
        report = await process_document(s, emb, settings, org_id=1, user_id=1,
                                        original_name="m2.xlsx", content=p.read_bytes())
        await s.commit()
    assert not any("не сходится" in w for w in report.warnings)

    await llm.close()
    await engine.dispose()


# ------------------------------------------------------------ «категория | значение» без периодов

def test_categorical_layout_without_periods(tmp_path):
    """Сводка клиента «за 2025 год сделки заключенные: источник | сумма» —
    периодов нет, но строки должны стать фактами, заголовок — родителем,
    итог — досчитанным."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["L3"] = "за 2025 год сделки заключенные"
    for i, (k, v) in enumerate([("сайт", 436350), ("2 гис", 2116040), ("авито", 276000)], start=5):
        ws.cell(i, 12, k)
        ws.cell(i, 13, v)
    path = tmp_path / "сделки.xlsx"
    wb.save(path)
    doc = load_excel(str(path))
    by_name = {f.metric_name: f for f in doc.facts}
    assert {"сайт", "2 гис", "авито", "сделки заключенные"} <= set(by_name)
    assert by_name["2 гис"].period.label == "2025" and by_name["2 гис"].section == "сделки заключенные"
    total = by_name["сделки заключенные"]
    assert total.is_total and total.value == 436350 + 2116040 + 276000
    assert doc.sections == ["сделки заключенные"]
    assert not doc.warnings  # год указан в заголовке — предупреждения нет


def test_categorical_layout_year_missing_warns(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "за год сделки заключенные"
    for i, (k, v) in enumerate([("сайт", 1), ("парсер", 2), ("авито", 3)], start=2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
    path = tmp_path / "s.xlsx"
    wb.save(path)
    doc = load_excel(str(path))
    assert len(doc.facts) == 4
    assert any("без периодов" in w for w in doc.warnings)


def test_rule_classifier_client_phrasing():
    from app.llm import _mock_classify
    assert _mock_classify("из чего состоят сделки заключенные?")["intent"] == "breakdown"
    assert _mock_classify("какие источники самые слабые?")["intent"] == "rank"
    assert _mock_classify("какой источник принёс больше всего сделок за год?")["intent"] == "rank"


def test_categorical_layout_year_from_file_name(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "сделки заключенные"
    for i, (k, v) in enumerate([("сайт", 1), ("парсер", 2), ("авито", 3)], start=2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
    path = tmp_path / "s.xlsx"
    wb.save(path)
    doc = load_excel(str(path), doc_name="сделки за 2024.xlsx")
    assert {f.period.label for f in doc.facts} == {"2024"} and not doc.warnings


async def test_compare_two_metrics(tmp_path):
    """«сравни авито и сайт» — два показателя за один период."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "за 2025 год сделки заключенные"
    for i, (k, v) in enumerate([("сайт", 436350), ("2 гис", 2116040), ("авито", 276000)], start=2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
    path = tmp_path / "сделки.xlsx"
    wb.save(path)
    from app.qa import AnswerPipeline

    settings.database_url = f"sqlite+aiosqlite:///{tmp_path / 'cmp.db'}"
    settings.llm_provider, settings.embeddings_provider, settings.send_charts = "mock", "hash", False
    settings.classify_with_llm = False
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    async with sessions() as s:
        await process_document(s, emb, settings, org_id=1, user_id=1,
                               original_name="сделки.xlsx", content=path.read_bytes())
        await s.commit()
    pipe = AnswerPipeline(sessions, emb, make_llm(settings), settings)
    out = await pipe.answer(1, 1, "сравни авито и сайт")
    assert "Сравнение «авито» и «сайт» за 2025" in out.text
    assert "меньше на" in out.text and "160,35 тыс." in out.text
    # обычное сравнение периодов не сломано: «и» внутри вопроса без второго показателя
    out = await pipe.answer(1, 1, "на сколько выросли сделки заключенные с 2024 по 2025?")
    assert "Сравнение «" not in out.text
    await engine.dispose()
