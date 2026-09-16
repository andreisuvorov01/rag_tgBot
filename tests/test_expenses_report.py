"""Отчёт «Личные расходы из бюджета компании»: файл с колонкой «Личные расходы».

Проверяем сквозной путь: сообщение в чат → журнал → отчёт в Excel, который
открывается и содержит ожидаемые суммы. Отдельно — точность: суммы в отчёте
обязаны сходиться копейка в копейку (деньги считаются Decimal).
"""
from __future__ import annotations

import asyncio
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl

from app.config import Settings
from app.expenses import add_entry, delete_last_entry, ledger_rows, parse_entry_message
from app.export import expenses_report_filename, expenses_report_xlsx
from app.storage import make_engine, make_sessionmaker

JOURNAL_NAME = "расходы.xlsx"


def _settings(tmp_path) -> Settings:
    # _env_file=None отключает чтение .env машины (в аннотациях pydantic-settings
    # такого параметра нет — отсюда точечное подавление для mypy)
    return Settings(  # type: ignore[call-arg]
        _env_file=None,
        database_url=f"sqlite+aiosqlite:///{tmp_path}/exp.db",
        data_dir=tmp_path,
        expense_journal_file=JOURNAL_NAME,
        llm_provider="mock",
    )


def _sheet_rows(data: bytes, title: str = "Отчёт") -> list[list]:
    wb = openpyxl.load_workbook(BytesIO(data))
    ws = wb[title]
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


async def _add(session, s: Settings, text: str, org_id: int = 1) -> dict:
    entry = parse_entry_message(text)
    assert entry is not None, f"не распознано: {text}"
    return await add_entry(session, s, org_id=org_id, user_id=1, entry=entry)


def test_order_of_words_does_not_matter():
    """«расход: кофе 1500» — описание перед суммой: это ожидаемая формулировка.

    Раньше парсер ждал сумму сразу после двоеточия, и такая запись не
    распознавалась вообще (сообщение уходило в QA-конвейер как вопрос).
    """
    e = parse_entry_message("расход: кофе 1500")
    assert e is not None
    assert e.kind == "expense"
    assert e.amount == Decimal("1500")
    assert "кофе" in e.description

    e2 = parse_entry_message("расход: 1500 кофе")
    assert e2 is not None and e2.amount == Decimal("1500")


def test_parser_handles_live_phrasings():
    """Формулировки, которые реально пишут в чат: порядок, валюта, дата."""
    today = date(2026, 1, 10)
    cases = {
        "расход: 1500 кофе": ("expense", Decimal("1500"), "кофе", today),
        "расход: кофе 1500": ("expense", Decimal("1500"), "кофе", today),
        "расход: 2 500,50 руб — такси": ("expense", Decimal("2500.50"), "такси", today),
        "расход: такси 700,50 руб": ("expense", Decimal("700.50"), "такси", today),
        "расходы - 1.500,00 ₽ кофе": ("expense", Decimal("1500.00"), "кофе", today),
        "расход: кофе с собой 120 р": ("expense", Decimal("120"), "кофе с собой", today),
        "расход: 300 обед (03.01)": ("expense", Decimal("300"), "обед", date(2026, 1, 3)),
        "расход: кофе 1500 (03.01.2026)": ("expense", Decimal("1500"), "кофе", date(2026, 1, 3)),
        "расход: 5 000 командировка (12.02)": ("expense", Decimal("5000"), "командировка",
                                               date(2026, 2, 12)),
        "расход: 1500": ("expense", Decimal("1500"), "", today),
        "доход: 50000 зарплата": ("income", Decimal("50000"), "зарплата", today),
    }
    for text, (kind, amount, desc, when) in cases.items():
        parsed = parse_entry_message(text, today=today)
        assert parsed is not None, f"не распознано: {text}"
        assert parsed.kind == kind, text
        assert parsed.amount == amount, text
        assert parsed.description == desc, f"{text}: {parsed.description!r}"
        assert parsed.when == when, text

    # обычный текст по-прежнему не считается записью расхода
    assert parse_entry_message("привет, как дела?", today=today) is None
    assert parse_entry_message("какая выручка за 2025?", today=today) is None


def test_report_contains_personal_expenses_column(tmp_path):
    """Главное требование: в отчёте есть колонка «Личные расходы» по месяцам."""
    s = _settings(tmp_path)

    async def run():
        engine = await make_engine(s)
        sessions = make_sessionmaker(engine)
        async with sessions() as session:
            await _add(session, s, "расход: кофе 1500 (03.01.2026)")
            await _add(session, s, "расход: такси 700 (03.01.2026)")
            await _add(session, s, "расход: обед 2300 (05.02.2026)")
            await _add(session, s, "доход: 50000 зарплата (05.02.2026)")
            await session.commit()
            rows = await ledger_rows(session, s, 1)
        await engine.dispose()
        return rows

    rows = asyncio.run(run())
    assert len(rows) == 4
    data = expenses_report_xlsx(rows, org_name="ООО Вектор", today=date(2026, 3, 1))
    table = _sheet_rows(data)

    # заголовок колонки и наличие обоих месяцев
    header = next(r for r in table if r and r[0] == "Месяц")
    assert header[:5] == ["Месяц", "Личные расходы", "Личные доходы", "Сальдо", "Записей"]
    labels = [r[0] for r in table if r and isinstance(r[0], str)]
    assert any("январь 2026" in x for x in labels)
    assert any("февраль 2026" in x for x in labels)

    jan = next(r for r in table if r and isinstance(r[0], str) and "январь 2026" in r[0])
    feb = next(r for r in table if r and isinstance(r[0], str) and "февраль 2026" in r[0])
    assert jan[1] == 2200.0        # 1500 + 700
    assert jan[2] == 0.0
    assert feb[1] == 2300.0
    assert feb[2] == 50000.0
    total = next(r for r in table if r and r[0] == "ИТОГО")
    assert total[1] == 4500.0
    assert total[2] == 50000.0
    # сальдо = расходы − доходы
    assert total[3] == 4500.0 - 50000.0


def test_report_has_categories_sheet(tmp_path):
    s = _settings(tmp_path)

    async def run():
        engine = await make_engine(s)
        sessions = make_sessionmaker(engine)
        async with sessions() as session:
            await _add(session, s, "расход: кофе 1500 (03.01.2026)")
            await _add(session, s, "расход: такси 700 (04.01.2026)")
            await session.commit()
            rows = await ledger_rows(session, s, 1)
        await engine.dispose()
        return rows

    rows = asyncio.run(run())
    data = expenses_report_xlsx(rows, today=date(2026, 2, 1))
    cats = _sheet_rows(data, "По категориям")
    assert cats[0][0] == "Категория"
    assert any("январь 2026" in str(c) for c in cats[0])
    body = [r for r in cats[1:] if r and r[0]]
    assert body, "должны быть строки категорий"
    # итоговая строка сходится с суммой расходов
    total_row = next(r for r in body if r[0] == "ИТОГО")
    assert total_row[-1] == 2200.0


def test_money_is_exact(tmp_path):
    """Копейки не должны расплываться: в отчёте ровно 0,1 + 0,2 = 0,3."""
    s = _settings(tmp_path)

    async def run():
        engine = await make_engine(s)
        sessions = make_sessionmaker(engine)
        async with sessions() as session:
            await _add(session, s, "расход: чай 0,10 (01.01.2026)")
            await _add(session, s, "расход: сок 0,20 (01.01.2026)")
            await session.commit()
            rows = await ledger_rows(session, s, 1)
        await engine.dispose()
        return rows

    rows = asyncio.run(run())
    assert sum(r["amount"] for r in rows) == Decimal("0.30")
    table = _sheet_rows(expenses_report_xlsx(rows, today=date(2026, 1, 2)))
    total = next(r for r in table if r and r[0] == "ИТОГО")
    assert abs(total[1] - 0.30) < 1e-9


def test_income_not_counted_as_expense(tmp_path):
    """Доходы идут отдельной колонкой и не попадают в «Личные расходы»."""
    s = _settings(tmp_path)

    async def run():
        engine = await make_engine(s)
        sessions = make_sessionmaker(engine)
        async with sessions() as session:
            await _add(session, s, "расход: кофе 100 (10.01.2026)")
            await _add(session, s, "доход: премия 9999 (10.01.2026)")
            await session.commit()
            rows = await ledger_rows(session, s, 1)
        await engine.dispose()
        return rows

    rows = asyncio.run(run())
    data = expenses_report_xlsx(rows, today=date(2026, 1, 11))
    table = _sheet_rows(data)
    jan = next(r for r in table if r and isinstance(r[0], str) and "январь 2026" in r[0])
    assert jan[1] == 100.0
    assert jan[2] == 9999.0
    # в структуру расходов доход не попал
    cat_names = [r[0] for r in table if r and isinstance(r[0], str) and r[0] not in ("Месяц", "ИТОГО")]
    assert not any("премия" in str(n).casefold() for n in cat_names)


def test_empty_journal_gives_empty_report():
    data = expenses_report_xlsx([], today=date(2026, 1, 1))
    table = _sheet_rows(data)
    assert any(r and r[0] == "Месяц" for r in table)
    # второго листа с данными нет — вместо него пояснение
    assert _sheet_rows(data, "По категориям")


def test_report_filename_has_month():
    assert expenses_report_filename(date(2026, 3, 14)) == "Личные_расходы_2026-03.xlsx"


def test_undo_removes_entry_from_report(tmp_path):
    """Отмена записи должна убирать её и из отчёта (он собирается из базы)."""
    s = _settings(tmp_path)

    async def run():
        engine = await make_engine(s)
        sessions = make_sessionmaker(engine)
        async with sessions() as session:
            await _add(session, s, "расход: кофе 1500 (03.01.2026)")
            await _add(session, s, "расход: такси 700 (04.01.2026)")
            await session.commit()
            before = await ledger_rows(session, s, 1)
            removed = await delete_last_entry(session, 1)
            await session.commit()
            after = await ledger_rows(session, s, 1)
        await engine.dispose()
        return before, removed, after

    before, removed, after = asyncio.run(run())
    assert len(before) == 2 and removed["amount"] == Decimal("700")
    assert len(after) == 1
    table = _sheet_rows(expenses_report_xlsx(after, today=date(2026, 1, 5)))
    total = next(r for r in table if r and r[0] == "ИТОГО")
    assert total[1] == 1500.0


def test_report_does_not_touch_uploaded_company_file(tmp_path):
    """Оригинал отчёта компании не меняется: пишем только новый файл."""
    s = _settings(tmp_path)
    original = tmp_path / "Отчет_компании.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Показатель", "2026"])
    wb.active.append(["Выручка", 1000])
    wb.save(original)
    before = original.read_bytes()

    async def run():
        engine = await make_engine(s)
        sessions = make_sessionmaker(engine)
        async with sessions() as session:
            await _add(session, s, "расход: кофе 1500 (03.01.2026)")
            await session.commit()
            rows = await ledger_rows(session, s, 1)
        await engine.dispose()
        return rows

    rows = asyncio.run(run())
    out = tmp_path / expenses_report_filename(date(2026, 1, 3))
    out.write_bytes(expenses_report_xlsx(rows, today=date(2026, 1, 3)))

    assert original.read_bytes() == before, "исходный файл компании изменён"
    assert out.exists() and out.stat().st_size > 0


def test_journal_file_still_written(tmp_path):
    """Прежний журнал расходы.xlsx продолжает обновляться — регрессия не допускается."""
    s = _settings(tmp_path)

    async def run():
        from app.expenses import sync_journal_file

        engine = await make_engine(s)
        sessions = make_sessionmaker(engine)
        async with sessions() as session:
            await _add(session, s, "расход: кофе 1500 (03.01.2026)")
            await session.commit()
            n = await sync_journal_file(session, s, 1)
        await engine.dispose()
        return n

    n = asyncio.run(run())
    assert n == 1
    path = s.expense_journal_path
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Дата", "Тип", "Сумма", "Категория", "Описание"]
    assert ws["A2"].value == "03.01.2026"
    assert ws["C2"].value == 1500
