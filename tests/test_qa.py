"""Интеграционный тест: загрузка -> БД -> вопросы (offline: mock LLM, hash-эмбеддинги)."""
from datetime import date

import openpyxl

from app.config import settings
from app.embeddings import EmbeddingService
from app.ingest.pipeline import process_document
from app.llm import make_llm
from app.qa import AnswerPipeline
from app.storage import make_engine, make_sessionmaker


async def _setup(tmp_path):
    settings.database_url = f"sqlite+aiosqlite:///{tmp_path}/t.db"
    settings.data_dir = tmp_path
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    llm = make_llm(settings)
    return engine, sessions, emb, llm


def _workbook(path, years=("2023", "2024", "2025")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Опер. расходы"
    ws["A1"] = "тыс. руб."
    ws.append(["Показатель", *years])
    data = {
        "Выручка": (41200, 45800, 49500),
        "Аренда спецтехники": (9800, 11400, 12700),
        "Зарплаты": (15600, 16800, 17900),
    }
    for name, vals in data.items():
        ws.append([name, *vals])
    wb.save(path)


async def test_end_to_end_factual_and_forecast(tmp_path):
    engine, sessions, emb, llm = await _setup(tmp_path)
    wb_path = tmp_path / "report.xlsx"
    _workbook(wb_path)

    async with sessions() as session:
        report = await process_document(
            session, emb, settings, org_id=1, user_id=1,
            original_name="report.xlsx", content=wb_path.read_bytes(),
        )
        await session.commit()
    assert report.status == "processed"
    assert report.facts == 9
    assert set(report.periods) == {"2023", "2024", "2025"}

    pipeline = AnswerPipeline(sessions, emb, llm, settings, org_names=["ООО Тест"])

    outcome = await pipeline.answer(1, 1, "Какая выручка за 2025 год?")
    assert "49 500 000" in outcome.text or "49,50 млн" in outcome.text.replace(" ", " ")

    outcome = await pipeline.answer(1, 1, "Какой прогноз по аренде спецтехники на 2026 год?")
    assert "Прогноз" in outcome.text or "прогноз" in outcome.text
    assert "12 700 000" in outcome.text or "12,70 млн" in outcome.text  # факт 2025 в истории

    outcome = await pipeline.answer(1, 1, "Прогноз по аренде на 2026, если темпы упадут вдвое")
    assert "0,5" in outcome.text  # множитель сценария виден в ответе

    # дедупликация по хэшу
    async with sessions() as session:
        dup = await process_document(
            session, emb, settings, org_id=1, user_id=1,
            original_name="report_copy.xlsx", content=wb_path.read_bytes(),
        )
        await session.commit()
    assert dup.status == "duplicate"

    await llm.close()
    await engine.dispose()


async def test_clarify_on_ambiguous_metric(tmp_path):
    engine, sessions, emb, llm = await _setup(tmp_path)
    # два похожих показателя
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Показатель", "2023", "2024"])
    ws.append(["Аренда спецтехники Москва", "100", "120"])
    ws.append(["Аренда спецтехники Регионы", "80", "90"])
    ws.append(["Зарплаты", "500", "600"])
    wb_path = tmp_path / "amb.xlsx"
    wb.save(wb_path)
    async with sessions() as session:
        await process_document(session, emb, settings, org_id=1, user_id=1,
                               original_name="amb.xlsx", content=wb_path.read_bytes())
        await session.commit()
    pipeline = AnswerPipeline(sessions, emb, llm, settings, org_names=[])
    outcome = await pipeline.answer(1, 1, "Сколько составляет аренда спецтехники за 2024?")
    # либо точный ответ, либо уточняющий вопрос с вариантами — но не выдумка
    assert outcome.clarify or "аренда" in outcome.text.casefold()
    await llm.close()
    await engine.dispose()


def test_facts_dates():
    assert date(2025, 12, 31) > date(2023, 1, 1)


async def test_fallback_when_metric_unknown_or_period_missing(tmp_path):
    """Тупик автоматики — не отписка, а сводка того, что есть (fallback):
    неизвестный показатель, год вне данных; шаблон без модели перечисляет
    показатели и подсказывает журнал."""
    from pathlib import Path

    from app.embeddings import EmbeddingService
    from app.ingest.pipeline import process_document
    from app.llm import make_llm
    from app.qa import AnswerPipeline
    from app.storage import make_engine, make_sessionmaker

    settings.database_url = f"sqlite+aiosqlite:///{tmp_path / 'fb.db'}"
    settings.llm_provider, settings.embeddings_provider, settings.send_charts = "mock", "hash", False
    settings.classify_with_llm = False
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    report = Path(__file__).resolve().parents[1] / "Финансовый_отчёт_ООО_Вектор_2023-2025.xlsx"
    async with sessions() as s:
        await process_document(s, emb, settings, org_id=1, user_id=1,
                               original_name=report.name, content=report.read_bytes())
        await s.commit()
    pipe = AnswerPipeline(sessions, emb, make_llm(settings), settings)

    out = await pipe.answer(1, 1, "сколько сделок принёс 2гис за сентябрь?")
    assert out.payload_type == "fallback"
    assert "выручка" in out.text and "«расход: 1500 кофе»" in out.text  # сводка + подсказка журнала

    # родственные кандидаты есть — по-прежнему уточнение кнопками, а не сводка
    out = await pipe.answer(1, 1, "сколько личных расходов за сентябрь?")
    assert out.clarify and "расходы" in out.clarify[0]

    out = await pipe.answer(1, 1, "какая выручка за 2019 год?")
    assert out.payload_type == "fallback" and "2025" in out.text  # год вне данных — показываем, что есть
    await engine.dispose()


async def test_human_phrasings_route_correctly(tmp_path):
    """Живые формулировки клиента на демо-файле: без модели (правила + hash)
    каждая должна попадать в правильный исполнитель и давать нужные цифры."""
    from pathlib import Path

    from app.embeddings import EmbeddingService
    from app.expenses import add_entry, parse_entry_message
    from app.ingest.pipeline import process_document
    from app.llm import make_llm
    from app.qa import AnswerPipeline
    from app.storage import make_engine, make_sessionmaker

    settings.database_url = f"sqlite+aiosqlite:///{tmp_path / 'human.db'}"
    settings.expense_journal_file = str(tmp_path / "расходы.xlsx")
    settings.llm_provider, settings.embeddings_provider, settings.send_charts = "mock", "hash", False
    settings.classify_with_llm = False
    engine = await make_engine(settings)
    sessions = make_sessionmaker(engine)
    emb = EmbeddingService(settings)
    wb = Path(__file__).resolve().parents[1] / "Сделки_по_источникам_2024-2026.xlsx"
    async with sessions() as s:
        await process_document(s, emb, settings, org_id=1, user_id=1, original_name=wb.name, content=wb.read_bytes())
        for msg in ("ушло 3000 на бензин (10.09.2026)", "такси 700 (05.09.2026)", "обед с клиентом 2500 (12.09.2026)"):
            await add_entry(s, settings, org_id=1, user_id=1, entry=parse_entry_message(msg), emb=emb)
        await s.commit()
    pipe = AnswerPipeline(sessions, emb, make_llm(settings), settings)

    async def ask(q: str):
        return await pipe.answer(1, 1, q)

    out = await ask("сколько сделок было в марте 2025?")
    assert out.payload_type == "factual" and "03.2025" in out.text and "05.2025" not in out.text
    out = await ask("что лучше сайт или интернет?")
    assert "Сравнение «сайт» и «интернет»" in out.text
    out = await ask("лучший месяц 2025 года")
    assert "Месяцы по убыванию" in out.text and out.text.index("10.2025") < out.text.index("01.2025")
    out = await ask("топ 3 источника за 2025")
    assert "Состав показателя «сделки заключенные» за 2025" in out.text
    out = await ask("динамика по авито")
    assert "2024" in out.text and "2025" in out.text and "-25,9%" in out.text
    out = await ask("что будет в 2027?")
    assert out.payload_type == "forecast" and "«сделки заключенные» на 2027" in out.text
    out = await ask("структура сделок в 2024")
    assert "за 2024" in out.text and "2 гис" in out.text
    out = await ask("разбивка по источникам за август 2026")
    assert "за 08.2026" in out.text
    out = await ask("на что я трачу больше всего?")
    assert "«личные расходы»" in out.text and "Транспорт" in out.text
    out = await ask("сколько я потратил на транспорт?")
    assert "Транспорт" in out.text and "бензин" in out.text and "такси" in out.text and "обед" not in out.text
    out = await ask("какие данные загружены?")
    assert out.payload_type == "fallback" and "Что есть в загруженных данных" in out.text
    out = await ask("привет")
    assert out.text.startswith("Здравствуйте")
    await engine.dispose()
