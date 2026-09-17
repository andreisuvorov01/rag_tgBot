"""Демо-файл по данным клиента: сделки по источникам за 2024–2026 помесячно.

Годовые суммы за 2025 совпадают с таблицей клиента (круговая диаграмма
«за год сделки заключенные»), 2024 и 2026 достроены с трендом по каждому
источнику, к суммам добавлено количество сделок. Два листа:
«Сделки, руб» и «Сделки, шт» — показатель × месяц, строка «Итого».

Запуск: python -m scripts.make_client_demo  -> Сделки_по_источникам_2024-2026.xlsx
"""
from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# источник -> (сумма за 2025 из таблицы клиента, 2024 к 2025, тренд 2026 к 2025, средний чек)
SOURCES: dict[str, tuple[int, float, float, int]] = {
    "сайт": (436_350, 0.90, 1.10, 12_000),
    "парсер": (1_013_500, 0.70, 1.25, 18_000),
    "интернет": (437_000, 1.05, 0.95, 14_000),
    "сарафанка": (1_123_390, 0.85, 1.15, 25_000),
    "2 гис": (2_116_040, 0.60, 1.30, 20_000),
    "авито": (276_000, 1.35, 0.80, 9_000),
    "повторник": (442_900, 0.80, 1.20, 30_000),
    "бартер": (160_590, 1.10, 0.90, 12_000),
    "бренд Климов": (1_045_670, 0.75, 1.20, 40_000),
    "зашли в офис": (301_650, 1.00, 1.00, 15_000),
    "соседи": (147_000, 1.20, 0.85, 8_000),
}
# сезонность: январь и май проседают, осень сильнее
SEASON = [0.70, 0.85, 1.00, 1.05, 0.80, 0.95, 1.00, 1.05, 1.20, 1.25, 1.15, 1.00]
MONTHS = ("январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь")
YEARS = ((2024, 12), (2025, 12), (2026, 9))  # 2026 — по сентябрь

_FILL = PatternFill("solid", fgColor="FFC000")


def _split_year(total: int, n_months: int, rng: random.Random) -> list[int]:
    """Годовую сумму — по месяцам с сезонностью и шумом, сумма сходится копейка в копейку."""
    weights = [SEASON[i] * rng.uniform(0.85, 1.15) for i in range(n_months)]
    scale = total / sum(weights)
    parts = [int(round(w * scale / 10) * 10) for w in weights]
    parts[-1] += total - sum(parts)
    return parts


def build(path: Path, seed: int = 7) -> None:
    rng = random.Random(seed)
    sums: dict[str, list[int]] = {}
    counts: dict[str, list[int]] = {}
    for name, (base_2025, k_2024, k_2026, check) in SOURCES.items():
        row_sum: list[int] = []
        for year, n in YEARS:
            total = {2024: round(base_2025 * k_2024), 2025: base_2025, 2026: round(base_2025 * k_2026 * n / 12)}[year]
            row_sum += _split_year(total, n, rng)
        sums[name] = row_sum
        counts[name] = [max(1, round(v / check)) if v > 0 else 0 for v in row_sum]

    headers = [f"{MONTHS[m]} {year}" for year, n in YEARS for m in range(n)]
    wb = Workbook()
    for title, section, data, fmt in (
        ("Сделки, руб", "Сделки заключенные", sums, "#,##0"),
        ("Сделки, шт", "Количество сделок", counts, "0"),
    ):
        ws = wb.active if title.endswith("руб") else wb.create_sheet()
        ws.title = title
        ws["A1"] = "Сделки по источникам, 2024–2026"
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(["Источник", *headers])
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        # строка-раздел под шапкой: источники становятся её детьми, «Итого» — её итогом
        ws.append([section])
        ws["A4"].font = Font(bold=True)
        ws["A4"].fill = _FILL
        for name, values in data.items():
            label = name if title.endswith("руб") else f"{name}, сделок"
            ws.append([label, *values])
        totals = [sum(col) for col in zip(*data.values(), strict=True)]
        ws.append(["Итого", *totals])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        ws.column_dimensions["A"].width = 22
        for c in range(2, len(headers) + 2):
            ws.column_dimensions[get_column_letter(c)].width = 12
            for r in range(5, ws.max_row + 1):
                ws.cell(r, c).number_format = fmt
        ws.freeze_panes = "B5"
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).resolve().parents[1] / "Сделки_по_источникам_2024-2026.xlsx"
    build(out)
    print(f"записан {out}")
